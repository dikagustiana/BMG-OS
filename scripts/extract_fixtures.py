#!/usr/bin/env python3
"""
Extract golden fixtures and master seed data from the workpapers in reference/ (PLAN v1 Fase 0 step 3).

    python3 -I scripts/extract_fixtures.py            # write fixtures/ (month-2026-09, month-2026-08, budget-v1, master)
    python3 -I scripts/extract_fixtures.py --check    # extract to a temp dir and fail if fixtures/ differs (CI)
    python3 -I scripts/extract_fixtures.py --seed-sql supabase/migrations/<ts>_master_seed.sql
    python3 -I scripts/extract_fixtures.py --seed-sql <file> --tables product_item   # seed of one master (Fase 2)

Rules:
- Every workbook is read twice with openpyxl: data_only=True (Excel's cached values) and data_only=False (formulas).
- Numbers are written as decimal strings without exponent, exactly as the double Excel stored (shortest round-trip
  repr), never rounded. Engine code reads them with decimal.js.
- Nothing here computes a business number: fixtures are copies of WP cells; the seed only reshapes WP master lists.
"""
from __future__ import annotations

import csv
import datetime as dt
import filecmp
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
from decimal import Decimal

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REF = os.path.join(ROOT, "reference")
SEP = "DI - WP Costing BMG Sep26_4.xlsx"
AGU = "DI - WP Costing BMG Agu26_3.xlsx"        # referenced by Sep26_4 FCC roll-forward (FCC!E196, C201, C202)
AGU_ASUMSI = "WP Costing BMG Agu26.xlsx"         # same numbers as Agu26_3 plus the ASUMSI sheet (PRD D7)
BP = "BMG_Business_Plan_-_Oct26_to_Dec27.xlsx"

# ----------------------------------------------------------------------------------------------- cell values


def num_str(v) -> str:
    """Exact decimal string of an Excel number (int or IEEE double), no exponent, no trailing zeros."""
    if isinstance(v, bool):
        raise TypeError("bool is not a number")
    if isinstance(v, int):
        return str(v)
    s = format(Decimal(repr(float(v))), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


def cell_type(c) -> str:
    """n number, s string, b bool, e error, d date."""
    v = c.value
    if c.data_type == "e" or (isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")) and v in
                              ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!")):
        return "e"
    if isinstance(v, bool):
        return "b"
    if isinstance(v, (int, float)):
        return "n"
    if isinstance(v, (dt.datetime, dt.date)):
        return "d"
    return "s"


def cell_value(c) -> str:
    v = c.value
    t = cell_type(c)
    if t == "n":
        return num_str(v)
    if t == "b":
        return "TRUE" if v else "FALSE"
    if t == "d":
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    return str(v)


def formula_text(fc) -> str:
    v = fc.value
    if v is None:
        return ""
    if hasattr(v, "text"):  # ArrayFormula
        return str(v.text)
    s = str(v)
    return s if s.startswith("=") else ""


class Book:
    def __init__(self, name: str):
        path = os.path.join(REF, name)
        if not os.path.exists(path):
            sys.exit(f"reference file missing: {path}")
        self.name = name
        self.path = path
        with open(path, "rb") as f:
            self.sha256 = hashlib.sha256(f.read()).hexdigest()
        self.v = openpyxl.load_workbook(path, data_only=True)
        self.f = openpyxl.load_workbook(path, data_only=False)

    def sheet(self, name: str):
        return self.v[name]

    def val(self, sheet: str, ref: str):
        return self.v[sheet][ref].value

    def num(self, sheet: str, ref: str) -> str:
        c = self.v[sheet][ref]
        if cell_type(c) != "n":
            raise ValueError(f"{self.name} {sheet}!{ref} is not a number: {c.value!r}")
        return num_str(c.value)

    def cells(self):
        """All non-empty cells in sheet order, row-major: (sheet, coordinate, type, value, formula)."""
        for ws in self.v.worksheets:
            wf = self.f[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    yield ws.title, c.coordinate, cell_type(c), cell_value(c), formula_text(wf[c.coordinate])


# ----------------------------------------------------------------------------------------------- writers


def write_json(path: str, obj) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write("\n")


def write_cells_csv(path: str, book: Book) -> int:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    n_num = 0
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["sheet", "cell", "type", "value", "formula"])
        for sheet, coord, t, v, fx in book.cells():
            w.writerow([sheet, coord, t, v, fx])
            n_num += t == "n"
    return n_num


# ----------------------------------------------------------------------------------------------- sections

HEADING = re.compile(r"^\s*(\d+(?:[a-z]|–\d+)?)\.\s|^\s*(LAMPIRAN [A-Z])\.")


def section_key(text: str) -> str | None:
    m = HEADING.match(text)
    if not m:
        return None
    k = m.group(1) or m.group(2)
    return k.lower().replace(" ", "_").replace("–", "-")


def sheet_sections(book: Book, sheet: str) -> dict:
    """Split a WP sheet into its numbered parts (column A headings '1.', '2a.', 'LAMPIRAN B.' …).

    Rows keep every non-empty cell: c = {col: value string}, n = numeric columns, e = error columns.
    Rows above the first heading form the 'header' part (status lines, budget month selectors).
    """
    ws = book.sheet(sheet)
    parts: dict = {}
    order: list = []
    cur = "header"
    parts[cur] = {"title": ws["A1"].value or sheet, "first_row": 1, "rows": []}
    order.append(cur)
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        k = section_key(a) if isinstance(a, str) else None
        if k:
            base, i = k, 2
            while k in parts:  # a sheet can restart numbering (e.g. '5b.' after '5.')
                k = f"{base}_{i}"
                i += 1
            cur = k
            parts[cur] = {"title": a.strip(), "first_row": r, "rows": []}
            order.append(cur)
        cells, nums, errs = {}, [], []
        for c in ws[r]:
            if c.value is None:
                continue
            col = re.sub(r"\d+", "", c.coordinate)
            t = cell_type(c)
            cells[col] = cell_value(c)
            if t == "n":
                nums.append(col)
            elif t == "e":
                errs.append(col)
        if cells:
            row = {"r": r, "c": cells}
            if nums:
                row["n"] = nums
            if errs:
                row["e"] = errs
            parts[cur]["rows"].append(row)
    for k in order:
        rows = parts[k]["rows"]
        parts[k]["last_row"] = rows[-1]["r"] if rows else parts[k]["first_row"]
    return {k: parts[k] for k in order}


def sheets_fixture(book: Book, sheets: list[str]) -> dict:
    return {
        "source": {"file": book.name, "sha256": book.sha256},
        "sheets": {s: sheet_sections(book, s) for s in sheets},
    }


# ----------------------------------------------------------------------------------------------- month fixtures

SEP_FILES = {
    "summary.json": ["RINGKASAN"],
    "pool.json": ["BIAYA BERSAMA"],
    "fcc.json": ["FCC"],
    "uu.json": ["UU"],
    "uo.json": ["UO"],
    "others.json": ["LAIN-LAIN", "CEM-CEM ORI", "CEM-CEM BBQ", "KOPERASI", "MCH BASRENG L0", "MCH BASRENG L5",
                    "MCH BASRENG L10", "MCH KERIPIK LV3", "MCH KERIPIK LV10", "MCH KERIPIK LV5", "UBI PUTIH ACE"],
    "recon.json": ["REKON & KONTROL", "REKON MATERIAL"],
    "inventory.json": ["REKON PERSEDIAAN"],
    "reasonableness.json": ["CEK KEWAJARAN"],
    "bridge.json": ["BRIDGE GP"],
}


def T(id_, label, sheet, cell, expected, tol, src):
    return {"id": id_, "label": label, "sheet": sheet, "cell": cell, "expected": expected, "tolerance": tol, "source": src}


# PRD v1 §8 acceptance numbers and the WP cell each one comes from. Values are copied from the PRD text
# (not from the WP) so the integrity test proves the fixture and the PRD agree.
PRD_TARGETS = [
    T("B2.pool_total", "Total pool konversi", "BIAYA BERSAMA", "C87", "1572766737.19", "1", "PRD §8 B2"),
    T("B2.pool_A", "Pool Grup A", "BIAYA BERSAMA", "C81", "469314556.75", "1", "PRD §8 B2"),
    T("B2.pool_KUPAS", "Pool Kupas", "BIAYA BERSAMA", "C82", "53598624.96", "1", "PRD §8 B2"),
    T("B2.pool_B", "Pool Grup B", "BIAYA BERSAMA", "C83", "419567213.29", "1", "PRD §8 B2"),
    T("B2.pool_C", "Pool Grup C", "BIAYA BERSAMA", "C84", "46737260", "1", "PRD §8 B2"),
    T("B2.pool_D", "Pool Grup D", "BIAYA BERSAMA", "C85", "313485769.27", "1", "PRD §8 B2"),
    T("B2.pool_E", "Pool Grup E", "BIAYA BERSAMA", "C86", "270063312.92", "1", "PRD §8 B2"),
    T("B2.pool_fixed", "Pool tetap", "BIAYA BERSAMA", "D87", "806327113.62", "1", "PRD §8 B2"),
    T("B2.alloc_FCC", "Alokasi ke FCC", "BIAYA BERSAMA", "C121", "1287026419.50", "1", "PRD §8 B2"),
    T("B2.alloc_UU", "Alokasi ke UU", "BIAYA BERSAMA", "D121", "5780271.39", "1", "PRD §8 B2"),
    T("B2.alloc_UO", "Alokasi ke UO", "BIAYA BERSAMA", "E121", "78965900.38", "1", "PRD §8 B2"),
    T("B2.alloc_LAIN", "Alokasi ke Lain-lain", "BIAYA BERSAMA", "F121", "200994145.92", "1", "PRD §8 B2"),
    T("B2.fcc_cost", "FCC biaya produksi", "FCC", "E66", "5562568465.96", "1", "PRD §8 B2"),
    T("B2.fcc_cost_per_kg", "FCC biaya produksi per kg", "FCC", "F66", "23478.381", "0.001", "PRD §8 B2"),
    T("B2.fcc_rate", "FCC biaya rata-rata tertimbang per kg", "FCC", "D76", "23928.4692", "0.0001", "PRD §8 B2"),
    T("B2.fcc_hpp", "FCC HPP", "FCC", "E104", "5106312148.26", "1", "PRD §8 B2"),
    T("B2.fcc_transfer", "FCC transfer ke Lain-lain", "FCC", "E91", "551835128.88", "1", "PRD §8 B2"),
    T("B2.uu_hpp", "UU HPP", "UU", "E98", "232706417.35", "1", "PRD §8 B2"),
    T("B2.uo_hpp", "UO HPP (PRD: 252.566.6xx)", "UO", "E94", "252566650", "50", "PRD §8 B2"),
    T("B2.others_cost", "Lain-lain biaya produksi", "LAIN-LAIN", "G68", "870978077.67", "1", "PRD §8 B2"),
    T("B2.cemcem_ori_cost", "Cem-cem ORI biaya produksi", "LAIN-LAIN", "G47", "318461224.94", "1", "PRD §8 B2"),
    T("B2.others_hpp", "HPP Lain-lain", "LAIN-LAIN", "G93", "871581725.61", "1", "PRD §8 B2"),
    T("B3.hpp_card", "HPP kartu (Rp jt)", "RINGKASAN", "B29", "6463.1669", "0.0001", "PRD §8 B3"),
    T("B3.hpp_gl", "HPP GL (Rp jt)", "RINGKASAN", "B34", "6375.4123", "0.0001", "PRD §8 B3"),
    T("B3.recon_material", "Rekon: varian material (Rp jt)", "REKON & KONTROL", "C38", "-121.1482", "0.0001", "PRD §8 B3"),
    T("B3.recon_inv_wp", "Rekon: kenaikan persediaan WP (Rp jt)", "REKON & KONTROL", "C39", "-326.5394", "0.0001", "PRD §8 B3"),
    T("B3.recon_inv_gl", "Rekon: − kenaikan persediaan GL (Rp jt)", "REKON & KONTROL", "C40", "409.3055", "0.0001", "PRD §8 B3"),
    T("B3.recon_kw2", "Rekon: KW 2 pada NRV (Rp jt)", "REKON & KONTROL", "C41", "-25.5969", "0.0001", "PRD §8 B3"),
    T("B3.recon_b8", "Rekon: − penyesuaian bagian 8 (Rp jt)", "REKON & KONTROL", "C42", "-23.7756", "0.0001", "PRD §8 B3"),
    T("B3.recon_idle", "Rekon: jam menganggur (Rp jt)", "REKON & KONTROL", "C43", "0", "0.0001", "PRD §8 B3"),
    T("B3.recon_residual", "Rekon: sisa (≤ Rp 0,001 jt)", "RINGKASAN", "B35", "0", "0.001", "PRD §8 B3"),
    T("B3.bs_1106_001_wp", "1106-001 WP (Rp jt)", "RINGKASAN", "B46", "826.0313", "0.0001", "PRD §8 B3"),
    T("B3.bs_1106_001_gl", "1106-001 GL (Rp jt)", "RINGKASAN", "C46", "947.1795", "0.0001", "PRD §8 B3"),
    T("B3.bs_1106_001_expl", "1106-001 dijelaskan (Rp jt)", "RINGKASAN", "F46", "-121.1482", "0.0001", "PRD §8 B3"),
    T("B3.bs_wip_wp", "WIP WP (Rp jt)", "RINGKASAN", "B47", "51.9846", "0.0001", "PRD §8 B3"),
    T("B3.bs_wip_gl", "WIP GL (Rp jt)", "RINGKASAN", "C47", "297.8507", "0.0001", "PRD §8 B3"),
    T("B3.bs_wip_expl", "WIP dijelaskan (Rp jt)", "RINGKASAN", "F47", "-245.8661", "0.0001", "PRD §8 B3"),
    T("B3.bs_fg_wp", "FG WP (Rp jt)", "RINGKASAN", "B48", "14.1713", "0.0001", "PRD §8 B3"),
    T("B3.bs_fg_gl", "FG GL (Rp jt)", "RINGKASAN", "C48", "286.5366", "0.0001", "PRD §8 B3"),
    T("B3.bs_fg_expl", "FG dijelaskan (Rp jt)", "RINGKASAN", "F48", "-272.3654", "0.0001", "PRD §8 B3"),
    T("B3.k2_wip_close", "REKON PERSEDIAAN K2 WIP akhir", "REKON PERSEDIAAN", "H58", "225691966.25", "1", "PRD §8 B3"),
    T("B3.b21_card", "B21 FG masuk kartu (kg)", "FCC", "B87", "221027", "0.01", "PRD §8 B3 / D4"),
    T("B3.b21_accurate", "B21 FG masuk Accurate (kg)", "FCC", "C206", "211979.5", "0.01", "PRD §8 B3 / D4"),
    T("B3.b22_card", "B22 transfer kartu (kg)", "FCC", "B91", "23115.5", "0.01", "PRD §8 B3 / D4"),
    T("B3.b22_accurate", "B22 transfer Accurate (kg)", "FCC", "C207", "14058", "0.01", "PRD §8 B3 / D4"),
    T("B4.bridge_aug_pp", "Margin GL Agu (pp)", "BRIDGE GP", "C117", "19.9219", "0.0001", "PRD §8 B4"),
    T("B4.bridge_sep_pp", "Margin GL Sep (pp)", "BRIDGE GP", "C135", "14.5657", "0.0001", "PRD §8 B4"),
    T("CTX.production_kg", "Produksi total (kg)", "RINGKASAN", "F7", "289759.42", "0.01", "PRD §1"),
    T("CTX.revenue_jt", "Pendapatan (Rp jt)", "RINGKASAN", "F11", "7462.3602", "0.0001", "PRD §1"),
    T("CTX.idle_hours", "Jam menganggur Unit I", "BIAYA BERSAMA", "C107", "3264.8168", "0.0001", "PRD §1"),
]


def extract_month_09(out: str, sep: Book) -> dict:
    d = os.path.join(out, "month-2026-09")
    n_num = write_cells_csv(os.path.join(d, "cells.csv"), sep)
    for fname, sheets in SEP_FILES.items():
        write_json(os.path.join(d, fname), sheets_fixture(sep, sheets))
    write_json(os.path.join(d, "prd_targets.json"), {"source": "docs/PRD-v1-bmg-platform.md §8", "targets": PRD_TARGETS})
    return {"numeric_cells": n_num}


def extract_month_08(out: str, agu: Book, sep: Book) -> dict:
    d = os.path.join(out, "month-2026-08")
    n_num = write_cells_csv(os.path.join(d, "cells.csv"), agu)
    write_json(os.path.join(d, "summary.json"), sheets_fixture(agu, ["RINGKASAN"]))
    write_json(os.path.join(d, "pool.json"), sheets_fixture(agu, ["BIAYA BERSAMA"]))
    # Bridge inputs for August as the September WP used them (BRIDGE GP section 1, column C = Agu26_6).
    ws = sep.sheet("BRIDGE GP")
    rows = []
    for r in range(7, 78):
        code = ws.cell(r, 1).value
        if not code:
            continue
        rows.append({
            "row": r, "code": code, "label": ws.cell(r, 2).value,
            "aug": num_str(ws.cell(r, 3).value), "sep": num_str(ws.cell(r, 4).value), "address": ws.cell(r, 5).value,
        })
    write_json(os.path.join(d, "bridge_input.json"), {
        "source": {"file": sep.name, "sha256": sep.sha256, "range": "BRIDGE GP!A7:E77",
                   "note": "Kolom C = nilai Agu dari 'DI - WP Costing BMG Agu26_6' menurut WP Sep26_4 (file Agu26_6 tidak ada di reference/)."},
        "unit": "Rp juta; kg", "inputs": rows,
    })
    return {"numeric_cells": n_num, "bridge_inputs": len(rows)}


# ----------------------------------------------------------------------------------------------- budget (BP v1)

PRODUCT_WORDS = [
    ("FCC", re.compile(r"flat chip cassava|\bfcc\b", re.I)),
    ("UU", re.compile(r"ubi ungu", re.I)),
    ("UO", re.compile(r"ubi orange", re.I)),
    ("LAIN", re.compile(r"lain-lain|\bothers?\b", re.I)),
]
POOL_IN_LABEL = re.compile(r"\bGrup ([A-E])\b")


def slug(s: str) -> str:
    s = s.lower().replace("&", " dan ").replace("%", " pct ")
    s = re.sub(r"[^a-z0-9.]+", "_", s)
    return re.sub(r"_+", "_", s).strip("_")


def product_of(text: str) -> str | None:
    for code, rx in PRODUCT_WORDS:
        if rx.search(text):
            return code
    return None


def extract_budget(out: str, bp: Book) -> dict:
    ws = bp.sheet("Business Plan - Model Details")
    months = []
    for c in range(10, 34):  # J..AG
        h = ws.cell(4, c).value
        if not isinstance(h, dt.datetime):
            raise ValueError(f"BP header {ws.cell(4, c).coordinate} is not a date: {h!r}")
        kind = "actual" if (ws.cell(3, c).value or "").strip().lower() == "actual" else "plan"
        months.append((c, f"{h.year:04d}-{h.month:02d}", kind))
    L1 = L2 = L3 = None
    lines, pool_map, seen = [], {}, {}
    for r in range(5, ws.max_row + 1):
        labels = {c: ws.cell(r, c).value for c in range(1, 9) if ws.cell(r, c).value not in (None, "")}
        unit = ws.cell(r, 9).value
        nums = [(c, m, k, ws.cell(r, c).value) for c, m, k in months
                if isinstance(ws.cell(r, c).value, (int, float)) and not isinstance(ws.cell(r, c).value, bool)]
        if 1 in labels:  # section: A = number, B = title
            L1 = f"{labels[1]} {labels.get(2, '')}".strip()
            L2 = L3 = None
            label = None
        elif 3 in labels:  # subsection (optionally coded in B)
            L2 = f"{labels[2]} {labels[3]}".strip() if 2 in labels else str(labels[3]).strip()
            L3 = None
            label = " / ".join(str(labels[c]).strip() for c in sorted(labels) if c > 3) or None
        elif 4 in labels:
            d = str(labels[4])
            rest = " / ".join(str(labels[c]).strip() for c in sorted(labels) if c > 4)
            if d.startswith(" ") and L3:  # indented child row (product split of the line above)
                label = f"{L3} / {d.strip()}"
            else:
                L3 = f"{d.strip()} / {rest}" if rest else d.strip()
                label = L3
        elif labels:
            label = " / ".join(str(labels[c]).strip() for c in sorted(labels))
        else:
            label = None
        if not nums:
            continue
        name = label or L3 or L2 or L1 or f"row {r}"
        path = []
        for p in (L1, L2, name):  # drop empty levels and a level that repeats its parent (data row on a header line)
            if p and (not path or p != path[-1]):
                path.append(p)
        component = ".".join(slug(p) for p in path[1:] or path) + (f".{slug(str(unit))}" if unit else "")
        product = product_of(name) if label else None
        base = f"{product or 'ALL'}|{component}"
        n = seen.get(base, 0) + 1
        seen[base] = n
        if n > 1:
            component = f"{component}#{n}"
            base = f"{product or 'ALL'}|{component}"
        u = str(unit or "").strip()
        uu = u.upper()
        role = "qty" if uu in ("KG", "KWH", "MMBTU", "JAM", "PEOPLE") else "price" if uu in ("IDR/KG", "IDR") else \
            "amount" if uu in ("IDR MIL", "IDR MIO") else None
        lines.append({
            "bp_row": r, "product": product, "component": component, "unit": u or None,
            # BP labels the money unit "IDR Mil"/"IDR Mio" but the cells hold full Rupiah (Q76 revenue Agu-26 =
            # 9.793.125.452 = WP BRIDGE GP C7 × 1.000.000), so for those rows amount = the cell value.
            "maps_to": role, "driver_code_base": base, "bp_label": " > ".join(path),
            "values": {m: num_str(v) for c, m, k, v in nums},
        })
        m_pool = POOL_IN_LABEL.search(name)
        if m_pool and (L2 or "").startswith(("4.1", "4.3")):  # pool composition (4.1 per account, 4.3 per pool)
            pool_map[component] = {"bp_row_code": component, "bp_label": " > ".join(path), "pool": m_pool.group(1),
                                   "fixed_pct": None, "bp_row": r}
    d = os.path.join(out, "budget-v1")
    write_json(os.path.join(d, "budget_line.json"), {
        "source": {"file": bp.name, "sha256": bp.sha256, "sheet": "Business Plan - Model Details"},
        "budget_version": {"name": "BP Okt-26 v1", "status": "locked",
                           "note": "D9: laba jual aset 60 jt/bulan hardcode; pajak Sep–Des 2026 kosong; Lain-lain tidak membayar FCC. Dibekukan as-is (Q4)."},
        "months": [{"month": f"{m}-01", "column_kind": k} for _, m, k in months],
        "expand": "one budget_line per (row, month): driver_code = driver_code_base + '|' + YYYY-MM; value = values[YYYY-MM]; "
                  "qty/price/amount = value when maps_to says so",
        "rows": lines,
    })
    write_json(os.path.join(d, "budget_pool_map.json"), {
        "source": {"file": bp.name, "sha256": bp.sha256, "rule": "draft: baris bagian 4.1 dan 4.3 yang labelnya memuat 'Grup A…E'"},
        "rows": sorted(pool_map.values(), key=lambda x: x["bp_row"]),
    })
    return {"budget_rows": len(lines), "budget_lines": sum(len(x["values"]) for x in lines), "budget_pool_map": len(pool_map)}


# ----------------------------------------------------------------------------------------------- master data

GROUP_CODES = {  # LAIN-LAIN row → stable code (name prefix asserted below)
    22: ("CEMCEM_ORI", "Cem-cem ORI"), 23: ("CEMCEM_BBQ", "Cem-cem BBQ"), 24: ("KOPERASI", "Koperasi"),
    25: ("MCH_BASRENG", "Maicih basreng"), 26: ("MCH_KERIPIK_LV3_10", "Maicih keripik (LV 3"), 27: ("MCH_KERIPIK_LV5", "Maicih keripik LV 5"),
    28: ("FCC_MAICIH", "FCC Maicih"), 29: ("FCC_PUTIH", "FCC Putih"), 30: ("KECIMPRING", "Kecimpring"),
    31: ("WHOLE_CHIPS_REMPAH", "Whole chips rempah"), 32: ("MAXI", "Maxi"), 33: ("PELLET_UU", "Pellet ubi ungu"),
    34: ("KERIPIK_UU_RETAIL", "Keripik ubi ungu retail"), 35: ("WHOLE_CHIPS_UU", "Whole chips ubi ungu"), 36: ("UMBIIS", "Umbiis"),
    37: ("UU_CHIPS", "Ubi ungu chips"), 38: ("PELLET_UO", "Pellet ubi orange"), 39: ("UO_CHIPS", "Ubi orange chips"),
    40: ("PELLET_LAIN", "Pellet lainnya"), 41: ("UBI_PUTIH_ACE", "Ubi putih ACE"), 42: ("WRP_MAKLON", "WRP"),
}


def main_line(note: str) -> str | None:
    n = note.lower()
    if "cek bahan" in n:  # ASUMSI no. 14: main material not confirmed yet
        return None
    if "maklon" in n:
        return "maklon"
    if "ubi ungu" in n:
        return "ubi_ungu"
    if "ubi orange" in n:
        return "ubi_orange"
    if "singkong" in n:
        return "singkong"
    return None


def material_category(name: str) -> str:
    n = name.lower()
    rules = [
        ("singkong kulit", "singkong_kulit"), ("singkong daging", "singkong_daging"), ("olein", "olein"), ("cng", "cng"),
        ("ammonium", "amonium"), ("amonium", "amonium"), ("hdpe", "hdpe"), ("karton", "karton"), ("basreng bumbu", "bahan_maklon"),
        ("bumbu", "bumbu"), ("plastik", "kemasan"), ("polyfilm", "kemasan"), ("label", "kemasan"), ("plakban", "kemasan"),
        ("tali rafia", "kemasan"), ("hot ink", "kemasan"), ("coding tape", "kemasan"), ("benang", "kemasan"),
        ("ubi ungu", "ubi_ungu"), ("ubi orange", "ubi_orange"), ("ubi putih", "ubi_putih"),
    ]
    for key, cat in rules:
        if key in n:
            return cat
    return "other"


ACCOUNT_RE = re.compile(r"\b\d{4}-\d{3}\b")


def extract_master(out: str, sep: Book, agu: Book, agu_asumsi: Book) -> dict:
    d = os.path.join(out, "master")
    src = f"{sep.name}"
    # products
    products = []
    for code, sheet, nrv in (("FCC", "FCC", sep.num("FCC", "C209")), ("UU", "UU", sep.num("LAIN-LAIN", "C7")), ("UO", "UO", sep.num("LAIN-LAIN", "C7"))):
        ws = sep.sheet(sheet)
        title = re.sub(r"^KARTU BIAYA PRODUKSI\s*[-–]\s*", "", str(ws["A1"].value)).split(",")[0].strip()
        fg = None
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            m = re.match(r"^3\. Barang jadi \((FG-\d+) ", a) if isinstance(a, str) else None
            if m:
                fg = m.group(1)
                break
        products.append({"code": code, "name": title, "fg_item_code": fg, "kw2_nrv_per_kg": nrv, "wip_completion_pct": "1",
                         "sort_order": len(products) + 1,
                         "source": f"{src} {sheet}!A1; NRV {'FCC!C209' if code == 'FCC' else 'LAIN-LAIN!C7 (KW 2 ubi)'}"})
    # others groups and SKUs
    ll = sep.sheet("LAIN-LAIN")
    groups, sheet_group = [], {}
    for r, (code, prefix) in GROUP_CODES.items():
        name = str(ll.cell(r, 1).value)
        if not name.startswith(prefix):
            raise ValueError(f"LAIN-LAIN!A{r} = {name!r}, expected prefix {prefix!r}")
        note = str(ll.cell(r, 2).value or "")
        link = str(ll.cell(r, 10).value or "")
        for sh in re.findall(r"(?:sheet |, )([A-Z][A-Z0-9 \-]+?)(?=,|$)", link.replace("Tautan ke sheet ", "sheet ")):
            sheet_group[sh.strip()] = code
        groups.append({"code": code, "name": name, "main_line": main_line(note), "main_line_note": note, "fg_unit": "kg",
                       "wp_row": r, "sort_order": r - 21, "source": f"{src} LAIN-LAIN!A{r}:B{r}"})
    skus = []
    for sh in ["CEM-CEM ORI", "CEM-CEM BBQ", "MCH BASRENG L0", "MCH BASRENG L5", "MCH BASRENG L10", "MCH KERIPIK LV3",
               "MCH KERIPIK LV10", "MCH KERIPIK LV5", "UBI PUTIH ACE"]:
        ws = sep.sheet(sh)
        head = str(ws["A45"].value)
        m = re.match(r"^3\. Barang jadi \((FG-\d+) (.*)\) dan HPP$", head)
        if not m:
            raise ValueError(f"{sh}!A45 = {head!r}")
        kg, units = ws["B13"].value or 0, ws["C13"].value or 0
        unit = ("bag" if kg != units else "kg") if kg and units else None
        skus.append({"code": m.group(1), "others_group_code": sheet_group[sh], "name": m.group(2), "fg_unit": unit,
                     "kg_per_unit": num_str(Decimal(num_str(kg)) / Decimal(num_str(units))) if kg and units else None,
                     "wp_sheet": sh, "source": f"{src} {sh}!A45, B13:C13"})
    for s in skus:  # SKUs without production this month take the unit of their group siblings
        if s["fg_unit"] is None:
            sib = [x for x in skus if x["others_group_code"] == s["others_group_code"] and x["fg_unit"]]
            s["fg_unit"] = sib[0]["fg_unit"] if sib else "kg"
    kop = sep.sheet("KOPERASI")
    mapping = {k.strip(): v for k, v in re.findall(r"([^;:=]+?) = (FG-\d+)", str(kop["A139"].value).split(":", 1)[1])}
    for c in range(2, 13):  # B..L
        name = str(kop.cell(6, c).value).strip()
        kg, units = kop.cell(7, c).value or 0, kop.cell(8, c).value or 0
        skus.append({"code": mapping[name], "others_group_code": "KOPERASI", "name": name, "fg_unit": "bag",
                     "kg_per_unit": num_str(Decimal(num_str(kg)) / Decimal(num_str(units))) if kg and units else None,
                     "wp_sheet": "KOPERASI", "source": f"{src} KOPERASI!{kop.cell(6, c).coordinate}, A139"})
    for i, s in enumerate(skus):
        s["sort_order"] = i + 1
    units_by_group = {}
    for s in skus:
        units_by_group.setdefault(s["others_group_code"], set()).add(s["fg_unit"])
    for g in groups:
        if "bag" in units_by_group.get(g["code"], set()):
            g["fg_unit"] = "bag"
    # materials (REKON MATERIAL A7:A60)
    rm = sep.sheet("REKON MATERIAL")
    materials = []
    for r in range(7, 61):
        name = str(rm.cell(r, 1).value).strip()
        o = str(rm.cell(r, 15).value or "")
        code = re.search(r"\b(?:BB|BP|GP|CP|WIP)-\d+\b", o)
        lines = {k: Decimal(num_str(rm.cell(r, c).value or 0)) for k, c in (("FCC", 4), ("UU", 5), ("UO", 6), ("LAIN", 7))}
        best = max(lines.items(), key=lambda kv: kv[1])
        materials.append({"code": code.group(0) if code else None, "name": name, "unit": rm.cell(r, 2).value,
                          "category": material_category(name), "line_default": best[0] if best[1] > 0 else None,
                          "sort_order": r - 6, "source": f"{src} REKON MATERIAL!A{r}:O{r}"})
    # payroll groups (BIAYA BERSAMA A28:D41)
    bb = sep.sheet("BIAYA BERSAMA")
    payroll = []
    kinds = {"Operator harian": "operator", "Support harian": "support", "Borongan kupas": "kupas", "Staf bulanan pabrik": "staf"}
    pools = {"Grup A": "A", "Grup B": "B", "Kupas": "KUPAS"}
    for r in range(28, 42):
        grp, dept, pool = bb.cell(r, 1).value, bb.cell(r, 2).value, bb.cell(r, 4).value
        kind = kinds[grp]
        unit = dept.split()[-1] if kind == "operator" else None
        code = "STAF_PABRIK" if kind == "staf" else f"{kind.upper()}_{slug(re.sub(r'^Kupas ', '', dept)).upper()}"
        payroll.append({"code": code, "name": grp, "dept": dept, "kind": kind, "pool": pools[pool], "unit": unit, "wp_row": r,
                        "sort_order": r - 27, "source": f"{src} BIAYA BERSAMA!A{r}:D{r}"})
    # account → pool (Sep-26 confirmed; Okt-26 draft)
    fixed_by_account = {}
    for r in range(56, 79):
        b, nature = str(bb.cell(r, 2).value), str(bb.cell(r, 4).value)
        for acc in ACCOUNT_RE.findall(b):
            fixed_by_account[acc] = nature == "Tetap"
    coa = []
    sep_pool = {"Grup A": "A", "Grup B": "B", "Grup C": "C", "Grup D": "D", "Grup E": "E"}
    new_accounts = {}
    for r in range(7, 23):
        acc, name, pool_txt, new = bb.cell(r, 1).value, bb.cell(r, 2).value, bb.cell(r, 4).value, str(bb.cell(r, 5).value or "")
        if pool_txt == "Dipecah (bagian 2)":
            pool = "THR" if "THR" in name else "BPJS" if "Jamsostek" in name else "PAYROLL_SPLIT"
            fixed = True if pool in ("THR", "BPJS") else None
        else:
            pool = sep_pool[pool_txt]
            fixed = fixed_by_account.get(acc)
        coa.append({"account": acc, "account_name": name, "effective_from": "2026-01-01", "effective_to": "2026-09-30", "pool": pool,
                    "fixed_flag": fixed, "confirmed": True, "note": f"WP BIAYA BERSAMA bagian 1: {pool_txt}",
                    "source": f"{src} BIAYA BERSAMA!A{r}:E{r}"})
        new_accounts[acc] = (new, pool, fixed, name)
    for r in range(144, 156):
        acc, desc, target, nature = bb.cell(r, 1).value, bb.cell(r, 2).value, bb.cell(r, 4).value, bb.cell(r, 5).value
        if acc.startswith("5101"):
            continue  # cut-off / reclass inside the 16 accounts: voucher tags, not an account mapping
        pool = "THR" if str(target).startswith("THR") else "OPEX_PABRIK"
        coa.append({"account": acc, "account_name": desc, "effective_from": "2026-01-01", "effective_to": "2026-09-30", "pool": pool,
                    "fixed_flag": nature == "Tetap", "confirmed": True,
                    "note": f"WP BIAYA BERSAMA bagian 8 → {target} (lewat gl_voucher_tag)", "source": f"{src} BIAYA BERSAMA!A{r}:F{r}"})
        new_accounts.setdefault(acc, (acc, pool, nature == "Tetap", desc))
    # Okt-26: new payroll accounts 5101-1XY. Hypothesis (draft, confirmed = false): X = kelompok in the order of
    # bagian 2 (1 operator → A, 2 support → B, 3 kupas → KUPAS, 4 staf → B); Y = 1 upah, 2 THR, 3 bonus, 4 BPJS.
    grp_pool = {"1": ("A", "operator"), "2": ("B", "support"), "3": ("KUPAS", "kupas"), "4": ("B", "staf")}
    comp = {"1": "upah", "2": "THR", "3": "bonus", "4": "BPJS"}
    okt = {}
    for acc, (new, pool, fixed, name) in new_accounts.items():
        targets = [a.strip() for a in new.split("/")]
        expanded = []
        for t in targets:
            expanded.append(t if re.fullmatch(r"\d{4}-\d{3}", t) else f"{targets[0][:5]}{t}")  # "5101-111/121" → 5101-121
        for a in expanded:
            m = re.fullmatch(r"5101-1([1-4])([1-4])", a)
            if m:
                p, who = grp_pool[m.group(1)]
                c = comp[m.group(2)]
                f = True if c != "upah" else who == "staf"
                okt[a] = {"pool": p, "fixed_flag": f, "account_name": f"{c} {who} (dari {acc})",
                          "note": f"DRAFT dari kolom E {acc} ({new}); pola 5101-1XY: X=kelompok bagian 2, Y=komponen — perlu konfirmasi Accurate"}
            elif a == "5101-150":
                okt[a] = {"pool": "A", "fixed_flag": True, "account_name": f"pesangon (dari {acc})", "note": f"DRAFT dari kolom E {acc}"}
            else:
                prev = okt.get(a)
                okt[a] = {"pool": pool, "fixed_flag": fixed, "account_name": name if not prev else f"{prev['account_name']} + {name}",
                          "note": f"DRAFT dari kolom E {acc}" if not prev else f"{prev['note']} + {acc}"}
    for a in [f"5101-20{i}" for i in range(1, 8)]:
        okt[a] = {"pool": "MATERIAL", "fixed_flag": False, "account_name": "varian/pemakaian per jenis bahan",
                  "note": "DRAFT: REKON & KONTROL!A34 — pemakaian per jenis bahan mulai Okt-26 (akun 5101-201 s/d 207)"}
    for a, nm in (("5101-301", "Varian Produksi Singkong"), ("5101-302", "Varian Produksi Ubi Ungu"), ("5101-303", "Varian Produksi Ubi Orange")):
        okt[a] = {"pool": "MATERIAL", "fixed_flag": False, "account_name": nm, "note": "DRAFT: REKON & KONTROL!A31:A33"}
    for a in sorted(okt):
        o = okt[a]
        coa.append({"account": a, "account_name": o["account_name"], "effective_from": "2026-10-01", "effective_to": None, "pool": o["pool"],
                    "fixed_flag": o["fixed_flag"], "confirmed": False, "note": o["note"], "source": f"{src} BIAYA BERSAMA!E7:E22 (akun baru mulai Okt-26)"})
    # assumptions: ASUMSI no. 1–14 (Agu26) + PRD §9 defaults
    asu = agu_asumsi.sheet("ASUMSI")
    keys = {1: "cost_system", 2: "payroll_split_basis", 3: "kupas_thr_bpjs", 4: "thr_bpjs_split_basis", 5: "nrv_kw2",
            6: "cng_m3_per_mmbtu", 7: "kw2_ubi_value", 8: "allocation_drivers", 9: "normal_capacity_labour",
            10: "wip_open_fcc_aug_basis", 11: "material_price_source", 12: "value_without_qty", 13: "others_produced_equals_sold",
            14: "others_main_line_unconfirmed"}
    nums = {3: sep.num("BIAYA BERSAMA", "C46"), 5: sep.num("FCC", "C209"), 6: sep.num("FCC", "C210"), 7: sep.num("LAIN-LAIN", "C7")}
    units = {3: "flag", 5: "Rp/kg", 6: "m3/MMBTU", 7: "Rp/kg"}
    assumptions = []
    for r in range(5, 19):
        no = asu.cell(r, 1).value
        assumptions.append({
            "key": keys[no], "value_num": nums.get(no), "value_text": str(asu.cell(r, 3).value), "unit": units.get(no),
            "effective_from": "2026-08-01", "owner": asu.cell(r, 5).value, "status": asu.cell(r, 6).value, "used_in": asu.cell(r, 8).value,
            "note": f"ASUMSI no. {no}: {asu.cell(r, 2).value}. Dasar: {asu.cell(r, 4).value}. Terdokumentasi: {asu.cell(r, 7).value}",
            "source": f"{agu_asumsi.name} ASUMSI!A{r}:H{r}" + (f"; nilai: {sep.name}" if no in nums else ""),
        })
    prd = [
        ("idle_to_period", sep.num("BIAYA BERSAMA", "C110"), None, "flag", "PRD D2/Q9: jam menganggur tetap diserap produk (0) sampai kapasitas normal ditetapkan", f"{sep.name} BIAYA BERSAMA!C110"),
        ("fg_target_days", "7", None, "hari", "PRD Q12: target stok FG untuk rencana = 1 minggu jual", "PRD v1 §9 Q12"),
        ("bom_band_pct", "0.15", None, "rasio", "PRD Q13: pita BOM untuk validasi bon ±15% dari standar", "PRD v1 §9 Q13"),
        ("variance_alert_pct", "0.02", None, "rasio", "PRD Q14: ambang alert varian minggu > 2% biaya standar minggu, atau 3 minggu tanda sama", "PRD v1 §9 Q14"),
        ("variance_alert_weeks", "3", None, "minggu", "PRD Q14: 3 minggu berturut tanda sama", "PRD v1 §9 Q14"),
        ("lock_deadline", None, "H+1 12:00; Senin 10:00 = BELUM DIISI", None, "PRD Q15: hari kunci", "PRD v1 §9 Q15"),
        ("week_start_dow", "1", "Senin", "ISO dow", "PRD Q20: minggu dimulai Senin; biaya tetap dipecah per hari kalender", "PRD v1 §9 Q20"),
        ("wip_open_over_production_ratio", "3", None, "x", "PRD D3/B34: WIP awal > 3× produksi bulan → PERIKSA", "PRD v1 §1 D3"),
    ]
    for key, vn, vt, unit, note, source in prd:
        assumptions.append({"key": key, "value_num": vn, "value_text": vt, "unit": unit, "effective_from": "2026-08-01", "owner": "Finance",
                            "status": "Default PRD", "used_in": None, "note": note, "source": source})
    product_items = extract_product_items(sep, products)
    data = {
        "product": products, "others_group": groups, "sku": skus, "material_item": materials, "payroll_group": payroll,
        "coa_pool_map": coa, "assumption": assumptions, "product_item": product_items,
    }
    for k, rows in data.items():
        write_json(os.path.join(d, f"{k}.json"), {"source": {"file": sep.name, "sha256": sep.sha256}, "rows": rows})
    return {k: len(v) for k, v in data.items()}


def extract_product_items(sep: Book, products: list[dict]) -> list[dict]:
    """Fase 2: product ↔ Accurate item code (roles fg, wip, sales, kw2_sales), read from the WP:
    fg = product sheet bagian 3 title; wip = REKON PERSEDIAAN K1 row whose source names the product sheet;
    sales / kw2_sales = the Sales By Item item named in LAMPIRAN A ("WIP-…" = the product's WIP item)."""
    src = sep.name
    rp = sep.sheet("REKON PERSEDIAAN")
    by_name, wip_of = {}, {}
    for r in range(55, rp.max_row + 1):
        code, name, note = rp.cell(r, 1).value, rp.cell(r, 2).value, str(rp.cell(r, 20).value or "")
        if not isinstance(code, str) or not code.startswith(("WIP-", "FG-")):
            continue
        by_name.setdefault(str(name).strip(), (code, r))
        m = re.search(r"WP: sheet (\w+)", note)
        if m and code.startswith("WIP-"):
            wip_of[m.group(1)] = (code, r)
    rows = []
    for p in products:
        sheet = p["code"]
        rows.append({"product": sheet, "role": "fg", "item_code": p["fg_item_code"], "source": f"{src} {sheet} bagian 3 (judul)"})
        code, r = wip_of[sheet]
        rows.append({"product": sheet, "role": "wip", "item_code": code, "source": f"{src} REKON PERSEDIAAN!A{r}, T{r}"})
        ws = sep.sheet(sheet)
        for r in range(1, ws.max_row + 1):
            label = str(ws.cell(r, 1).value or "")
            role = "sales" if label.startswith("Penjualan") else "kw2_sales" if label.startswith("KW 2 tawar terjual") else None
            if not role:
                continue
            m = re.search(r"Sales By Item:\s*(.+?)\s*(?:\(|$)", str(ws.cell(r, 5).value or ""))
            if not m:
                raise ValueError(f"{sheet}!E{r}: Sales By Item item not found")
            name = m.group(1).strip()
            if name.startswith("WIP-"):
                code, rr = wip_of[sheet]
                where = f"{sheet}!E{r} ({name} = WIP produk), REKON PERSEDIAAN!A{rr}"
            else:
                code, rr = by_name[name]
                where = f"{sheet}!E{r}, REKON PERSEDIAAN!A{rr}:B{rr}"
            rows.append({"product": sheet, "role": role, "item_code": code, "source": f"{src} {where}"})
    return rows


# ----------------------------------------------------------------------------------------------- seed SQL

SEED_COLUMNS = {
    "product": ["code", "name", "fg_item_code", "kw2_nrv_per_kg", "wip_completion_pct", "sort_order", "source"],
    "others_group": ["code", "name", "main_line", "main_line_note", "fg_unit", "wp_row", "sort_order", "source"],
    "sku": ["code", "others_group_code", "name", "fg_unit", "kg_per_unit", "wp_sheet", "sort_order", "source"],
    "material_item": ["code", "name", "unit", "category", "line_default", "sort_order", "source"],
    "payroll_group": ["code", "name", "dept", "kind", "pool", "unit", "wp_row", "sort_order", "source"],
    "coa_pool_map": ["account", "account_name", "effective_from", "effective_to", "pool", "fixed_flag", "confirmed", "note", "source"],
    "assumption": ["key", "value_num", "value_text", "unit", "effective_from", "owner", "status", "used_in", "note", "source"],
    "product_item": ["product", "role", "item_code", "source"],
}
CONFLICT = {"product": "(code)", "others_group": "(code)", "sku": "(code)", "material_item": "(name)", "payroll_group": "(code)",
            "coa_pool_map": "(account, effective_from)", "assumption": "(key, effective_from)", "product_item": "(product, role, item_code)"}


def sql_lit(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def seed_sql(master_dir: str, tables: list[str] | None = None) -> str:
    out = io.StringIO()
    out.write("-- BMG-OS master seed (PLAN v1 Fase 0 step 2). Generated by `python3 -I scripts/extract_fixtures.py --seed-sql`\n")
    out.write("-- from fixtures/master/*.json (extracted from reference/ workpapers). Insert-only: on conflict do nothing.\n")
    out.write("-- Changing master data later = a new migration or the app (audited), never an edit of this file.\n")
    for table in tables or ["product", "others_group", "sku", "material_item", "payroll_group", "coa_pool_map", "assumption"]:
        with open(os.path.join(master_dir, f"{table}.json"), encoding="utf-8") as f:
            rows = json.load(f)["rows"]
        cols = SEED_COLUMNS[table]
        out.write(f"\n-- {table}: {len(rows)} rows\n")
        out.write(f"insert into public.{table} ({', '.join(cols)}) values\n")
        out.write(",\n".join("  (" + ", ".join(sql_lit(r.get(c)) for c in cols) + ")" for r in rows))
        out.write(f"\non conflict {CONFLICT[table]} do nothing;\n")
    if tables:
        return out.getvalue()
    out.write("""
-- budget version v1 (PRD Q4): frozen as-is; lines are loaded in Fase 1 through the budget_bp upload.
insert into public.budget_version (name, status, note, source) values
  ('BP Okt-26 v1', 'locked',
   'D9: laba jual aset 60 jt/bulan hardcode; pajak Sep–Des 2026 kosong; Lain-lain tidak membayar FCC. Dibekukan as-is (Q4); perbaikan = v2 oleh manusia.',
   'reference/BMG_Business_Plan_-_Oct26_to_Dec27.xlsx')
on conflict (name) do nothing;
""")
    return out.getvalue()


# ----------------------------------------------------------------------------------------------- main


def extract_all(out: str) -> dict:
    sep, agu, agu_asumsi, bp = Book(SEP), Book(AGU), Book(AGU_ASUMSI), Book(BP)
    report = {"month-2026-09": extract_month_09(out, sep), "month-2026-08": extract_month_08(out, agu, sep),
              "budget-v1": extract_budget(out, bp), "master": extract_master(out, sep, agu, agu_asumsi)}
    # roll-forward sanity: Sep opening FCC WIP value = Agu26_3 closing (Sep FCC!C196 cites Agu26_3 FCC!E78)
    # (the Sep cell holds the pasted value, so compare at Rp 0,01 rather than bit-for-bit)
    if abs(Decimal(agu.num("FCC", "E78")) - Decimal(sep.num("FCC", "C196"))) > Decimal("0.01"):
        raise ValueError(f"Roll-forward mismatch: Agu26_3 FCC!E78 {agu.num('FCC', 'E78')} vs Sep FCC!C196 {sep.num('FCC', 'C196')}")
    sources = {b.name: b.sha256 for b in (sep, agu, agu_asumsi, bp)}
    write_json(os.path.join(out, "SOURCES.json"), {"reference": sources, "extractor": "scripts/extract_fixtures.py"})
    return report


def dirs_equal(a: str, b: str) -> list[str]:
    diffs = []
    cmp = filecmp.dircmp(a, b)

    def walk(c, prefix=""):
        diffs.extend(prefix + x for x in c.left_only + c.right_only + c.diff_files)
        for name, sub in c.subdirs.items():
            walk(sub, prefix + name + "/")

    walk(cmp)
    return diffs


def main(argv: list[str]) -> int:
    if "--seed-sql" in argv:
        path = argv[argv.index("--seed-sql") + 1]
        tables = argv[argv.index("--tables") + 1].split(",") if "--tables" in argv else None
        sql = seed_sql(os.path.join(ROOT, "fixtures", "master"), tables)
        with open(path, "w", encoding="utf-8", newline="\n") as f:
            f.write(sql)
        print(f"wrote {path}")
        return 0
    if "--check" in argv:
        tmp = tempfile.mkdtemp()
        try:
            extract_all(tmp)
            diffs = dirs_equal(os.path.join(ROOT, "fixtures"), tmp)
            # hand-written READMEs and fixtures of later phases (uploads/, week-…/) are not produced by this script
            diffs = [d for d in diffs if not d.startswith(("uploads", "week-")) and not d.endswith("README.md")]
            if diffs:
                print("fixtures differ from a fresh extraction:\n  " + "\n  ".join(diffs))
                return 1
            print("fixtures match a fresh extraction of reference/")
            return 0
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    report = extract_all(os.path.join(ROOT, "fixtures"))
    print(json.dumps(report, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
